"""docx 그리기 — 계획을 python-docx 로. Word 와 구글 문서 둘 다에서 열리게 표·문단·인라인 그림만 쓴다
(떠 있는 도형 없음). 서식 값은 `worksheet.kit_style.resolve_kit` 이 킷 번호를 푼 실제 값이다."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

from worksheet.kit import Kit
from worksheet.plan import (
    CELL_ROLES, BandPlan, CellPlan, FigurePlan, HeadingPlan, KeywordPagePlan, ParagraphPlan, Plan, TablePlan,
    column_widths,
)
from worksheet.style import BoxStyle, ResolvedKit, Side

# 칸 안쪽 여백 — hwpx 표와 같은 값(python-hwpx 의 새 표 칸 여백: 좌우 510·상하 141 HWPUNIT)을 써야
# 칸 안 그림의 "칸 폭 − 좌우 여백" 계산이 두 형식에서 같다.
_칸_여백 = {"left": 510, "right": 510, "top": 141, "bottom": 141}
# 연속한 두 표 사이에 두는 구분 문단의 높이(pt) — docx 는 붙은 표를 하나로 합쳐 읽는다.
_구분_pt = 4

_정렬 = {
    "LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
    "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT, "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "DISTRIBUTE": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
}
_선 = {"SOLID": "single", "DASH": "dashed", "DOT": "dotted", "DOUBLE": "double", "NONE": "nil"}


# 반전 원문자(킷의 circle.fill 글자색) — 떠 있는 타원 대신 한 글자로 번호를 보인다. ⓴ 다음은 괄호 숫자.
원문자_글자 = "❶❷❸❹❺❻❼❽❾❿⓫⓬⓭⓮⓯⓰⓱⓲⓳⓴"


def 원문자(n: int) -> str:
    return 원문자_글자[n - 1] if 1 <= n <= len(원문자_글자) else f"({n})"


def _선폭_mm(line_width: str) -> float:
    # 한글 도형의 선 굵기(lineWidth)는 HWPUNIT 이다
    return int(line_width) / 7200 * 25.4


def twip(hwp: int) -> int:
    return round(hwp / 5)


def emu(hwp: int) -> int:
    return hwp * 127


def _sz(width_mm: float) -> int:
    return max(2, round(width_mm / 25.4 * 72 * 8))  # 1/8 pt


# Word 는 속성 요소의 **순서**를 스키마대로 요구한다(어기면 "읽을 수 없는 내용"으로 복구를 묻는다).
# 새 요소는 아래 순서에서 자기보다 뒤에 오는 첫 요소 앞에 끼운다.
_TBLPR = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
          "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
          "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
_TCPR = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
         "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_변_순서 = ("w:top", "w:left", "w:bottom", "w:right")
# python-docx 의 CT_PPr 은 w:rPr(문단부호 서식)을 등록 속성으로 안 두어(`get_or_add_rPr()`가 없다)
# _두기 로 직접 끼운다 — 그때 쓰는 pPr 자식 순서.
_PPR = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
        "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
        "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
        "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_RPR = ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
        "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
        "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz",
        "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
        "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _두기(부모, 태그: str, 속성: dict[str, str], 순서: tuple[str, ...] = ()):
    기존 = 부모.find(qn(태그))
    if 기존 is not None:
        부모.remove(기존)
    e = OxmlElement(태그)
    for k, v in 속성.items():
        e.set(qn(k), v)
    # 부모가 python-docx 의 등록 클래스가 아닐 수 있어(tblCellMar·tcBorders 는 맨 lxml 요소다)
    # `insert_element_before` 대신 직접 찾는다. 뒤에 올 요소가 없으면 끝에 붙는다.
    assert not 순서 or 태그 in 순서, f"{태그} 가 순서 목록에 없다"
    뒤 = {qn(t) for t in 순서[순서.index(태그) + 1:]} if 순서 else set()
    자리 = next((c for c in 부모 if c.tag in 뒤), None)
    if 자리 is None:
        부모.append(e)
    else:
        자리.addprevious(e)
    return e


def _글꼴_박기(run, 글꼴: str) -> None:
    rpr = run._r.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts")
        rpr.insert(0, f)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        f.set(qn(k), 글꼴)


def _그림_넣기(run, 경로: Path, 폭: int, 높이: int) -> None:
    """인라인 그림 — 둘레 여백(dist*)을 0으로 박는다. python-docx 는 이 속성을 쓰지 않는데, 빠지면
    LibreOffice 등이 기본 여백(약 9pt)을 넣어 그림이 오른쪽으로 밀리고 칸 밖으로 넘친다."""
    run.add_picture(str(경로), width=Emu(폭), height=Emu(높이))
    inline = run._r.find(f".//{qn('wp:inline')}")
    for k in ("distT", "distB", "distL", "distR"):
        inline.set(k, "0")


class DocxWriter:
    def __init__(self, kit: Kit, resolved: ResolvedKit) -> None:
        self.kit, self.r = kit, resolved
        self.doc = Document()
        self._마지막이_표 = False
        s = self.doc.sections[0]
        p = resolved.page
        s.page_width, s.page_height = Twips(twip(p.width)), Twips(twip(p.height))
        s.left_margin, s.right_margin = Twips(twip(p.left)), Twips(twip(p.right))
        # 한글은 머리말·꼬리말 띠를 위·아래 여백 **안쪽에 더해** 본문을 그 아래에서 시작한다. Word 의
        # 위 여백은 용지 끝에서 본문까지이므로 둘을 더해야 본문 자리가 같고, 머리말 글은 여백 끝에 놓인다.
        s.top_margin, s.bottom_margin = Twips(twip(p.top + p.header)), Twips(twip(p.bottom + p.footer))
        s.header_distance, s.footer_distance = Twips(twip(p.top)), Twips(twip(p.bottom))
        normal = self.doc.styles["Normal"]
        normal.font.name = kit.docx_font
        rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(k), kit.docx_font)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        # 서식을 따로 안 받는 문단(빈 칸·빈 문단 표시)도 틀의 11pt·1.15배가 아니라 본문 글자 크기·한 줄로
        normal.font.size = Pt(resolved.char["body"].size_pt)
        normal.paragraph_format.line_spacing = 1.0
        # 한국어 줄 나눔 규칙을 쓰게 동아시아 언어를 한국어로
        _두기(normal.element.get_or_add_rPr(), "w:lang", {"w:eastAsia": "ko-KR"}, _RPR)
        # Word 2013 이후 방식(호환 모드 15): 표 테두리를 들여쓰기 자리에 둔다. 14 이하에서는 첫 칸의
        # 글자를 맞추느라 표가 칸 여백만큼 왼쪽으로 나가 강조박스 안 상자가 한쪽으로 쏠린다.
        호환 = self.doc.settings.element.find(qn("w:compat"))  # python-docx 기본 틀에 늘 있다
        모드 = next((c for c in 호환.findall(qn("w:compatSetting")) if c.get(qn("w:name")) == "compatibilityMode"), None)
        if 모드 is None:
            모드 = OxmlElement("w:compatSetting")
            모드.set(qn("w:name"), "compatibilityMode")
            모드.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
            호환.append(모드)  # compatSetting 은 compat 의 마지막 자식 종류다
        모드.set(qn("w:val"), "15")

    # --- 글자·문단 -----------------------------------------------------------------------
    def _run(self, 문단, 글: str, 글자_역할: str):
        cs = self.r.char[글자_역할]
        run = 문단.add_run(글)
        run.font.size = Pt(cs.size_pt)
        run.bold = cs.bold
        run.underline = cs.underline
        run.font.color.rgb = RGBColor.from_string(cs.color.lstrip("#").upper())
        _글꼴_박기(run, self.kit.docx_font)
        return run

    def _문단_서식(self, 문단, 문단_역할: str, 글자_역할: str) -> None:
        ps = self.r.para[문단_역할]
        문단.alignment = _정렬.get(ps.align, WD_ALIGN_PARAGRAPH.LEFT)
        pf = 문단.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(0)
        if ps.line_percent:
            # 한글의 퍼센트 줄간격은 "글자 크기 × 비율"이다. Word 의 배수 줄간격은 글꼴마다 기준이
            # 달라 같은 모양이 안 나오므로, 같은 높이를 최소값(atLeast)으로 준다 — 글자가 잘리지 않는다.
            pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            pf.line_spacing = Pt(self.r.char[글자_역할].size_pt * ps.line_percent / 100)

    def _글_채우기(self, 문단, 글: str, 글자_역할: str) -> None:
        # 빈 칸에는 run 을 만들지 않는다 — 칸 높이는 행 최소 높이가 지킨다.
        # 한 칸 안의 줄바꿈은 한 문단 안의 줄 나눔이다(hwpx 셀과 같다)
        if not 글:
            return
        줄들 = 글.split("\n")
        for i, 줄 in enumerate(줄들):
            run = self._run(문단, 줄, 글자_역할)
            if i < len(줄들) - 1:
                run.add_break()

    def _구분(self) -> None:
        if self._마지막이_표:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before, pf.space_after = Pt(0), Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(_구분_pt)
            self._마지막이_표 = False

    # --- 표 ------------------------------------------------------------------------------
    def _음영(self, tcPr, 색: str | None) -> None:
        if 색:
            _두기(tcPr, "w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": 색.lstrip("#").upper()}, _TCPR)

    def _테두리(self, tcPr, box: BoxStyle | None) -> None:
        변들 = _두기(tcPr, "w:tcBorders", {}, _TCPR)
        for 이름 in ("top", "left", "bottom", "right"):
            side: Side | None = getattr(box, 이름) if box is not None else None
            if side is None or side.kind == "NONE":
                _두기(변들, f"w:{이름}", {"w:val": "nil"}, _변_순서)
            else:
                _두기(변들, f"w:{이름}", {
                    "w:val": _선.get(side.kind, "single"), "w:sz": str(_sz(side.width_mm)),
                    "w:color": side.color.lstrip("#").upper(), "w:space": "0",
                }, _변_순서)

    def _새_표(self, 컨테이너, 열폭: list[int], 행수: int):
        표 = 컨테이너.add_table(rows=행수, cols=len(열폭))
        tblPr = 표._tbl.tblPr
        _두기(tblPr, "w:tblW", {"w:w": str(sum(twip(w) for w in 열폭)), "w:type": "dxa"}, _TBLPR)
        _두기(tblPr, "w:tblLayout", {"w:type": "fixed"}, _TBLPR)
        마진 = _두기(tblPr, "w:tblCellMar", {}, _TBLPR)
        for 변, 값 in (("top", 141), ("left", 510), ("bottom", 141), ("right", 510)):
            _두기(마진, f"w:{변}", {"w:w": str(twip(값)), "w:type": "dxa"}, _변_순서)
        for gc, w in zip(표._tbl.tblGrid.findall(qn("w:gridCol")), 열폭):
            gc.set(qn("w:w"), str(twip(w)))
        for 행 in 표.rows:
            for 칸, w in zip(행.cells, 열폭):
                칸.width = Twips(twip(w))
        return 표

    def _칸(self, 칸, plan: CellPlan, 표_테두리: str, 폭: int) -> None:
        보더_키, 파라_키, 캐릭_키, _ = CELL_ROLES[plan.role]
        box = self.r.box[보더_키 or 표_테두리]
        tcPr = 칸._tc.get_or_add_tcPr()
        self._테두리(tcPr, box)
        self._음영(tcPr, box.fill)
        칸.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        문단 = 칸.paragraphs[0]
        if plan.picture is not None:
            self._문단_서식(문단, "center", 캐릭_키)
            run = 문단.add_run()
            _글꼴_박기(run, self.kit.docx_font)
            # 칸 폭은 twip 으로 반올림되지만 그림 EMU(HWPUNIT×127)는 정확해 반올림한 칸 안쪽을 1twip
            # 못 되게 넘을 수 있다 — 그림을 온 twip 으로 내려 잡고 칸 안쪽으로 막는다(비율 유지).
            안쪽 = 칸.width.twips - 2 * twip(_칸_여백["left"])
            폭_tw = min(plan.picture.width // 5, 안쪽)
            높이_tw = plan.picture.height * 폭_tw // plan.picture.width
            _그림_넣기(run, plan.picture.path, 폭_tw * 635, 높이_tw * 635)
        else:
            self._문단_서식(문단, 파라_키, 캐릭_키)
            self._글_채우기(문단, plan.text, 캐릭_키)
        if plan.nested is not None:
            assert len(plan.nested.rows) == 1 and len(plan.nested.rows[0]) == 1, "중첩표는 1×1(강조박스)만 그린다"
            안쪽폭 = column_widths(plan.nested, self.kit, 폭=폭 - _칸_여백["left"] - _칸_여백["right"])
            안쪽 = self._새_표(칸, 안쪽폭, len(plan.nested.rows))
            self._칸들_채우기(안쪽, plan.nested, 안쪽폭)
            # python-docx 의 cell.add_table 은 [기존 빈 문단][표][빈 문단]을 만든다. 구글 문서는 칸이
            # 표로 시작하는 것을 못 받아 전체 줄 높이의 빈 줄을 끼워 넣는다(실측 3쪽 → 2쪽) — 앞
            # 문단을 지우는 대신 1pt 로 눌러 남긴다(뒤 문단은 Word 가 칸 끝에 요구한다 — 마찬가지로
            # 최소 높이로. 문단부호 서식(rPr)의 글자 크기까지 눌러야 구글 문서도 한 줄로 접는다).
            첫 = 칸.paragraphs[0].paragraph_format
            첫.space_before, 첫.space_after = Pt(0), Pt(0)
            첫.line_spacing_rule, 첫.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(1)
            첫_rpr = _두기(칸.paragraphs[0]._p.get_or_add_pPr(), "w:rPr", {}, _PPR)
            _두기(첫_rpr, "w:sz", {"w:val": "2"}, _RPR)
            끝 = 칸.paragraphs[-1].paragraph_format
            끝.line_spacing_rule, 끝.line_spacing = WD_LINE_SPACING.EXACTLY, Pt(1)

    def _칸들_채우기(self, 표, plan: TablePlan, 열폭: list[int]) -> None:
        for 행, 행_plan in zip(표.rows, plan.rows):
            높이 = max(
                (칸.height if 칸.height is not None else self.kit.furniture["rowHeight"][CELL_ROLES[칸.role][3]])
                for 칸 in 행_plan
            )
            행.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            행.height = Twips(twip(높이))
            for c, (칸, 칸_plan) in enumerate(zip(행.cells, 행_plan)):
                폭 = 칸_plan.width if 칸_plan.width is not None else 열폭[c]
                self._칸(칸, 칸_plan, plan.border, 폭)

    def _표(self, plan: TablePlan) -> None:
        self._구분()
        열폭 = column_widths(plan, self.kit)
        표 = self._새_표(self.doc, 열폭, len(plan.rows))
        self._칸들_채우기(표, plan, 열폭)
        self._마지막이_표 = True

    # --- 문단·그림 ---------------------------------------------------------------------------
    def _문단(self, plan: ParagraphPlan) -> None:
        p = self.doc.add_paragraph()
        self._문단_서식(p, plan.para, plan.char)
        self._글_채우기(p, plan.text, plan.char)
        self._마지막이_표 = False

    def _그림(self, plan: FigurePlan) -> None:
        p = self.doc.add_paragraph()
        self._문단_서식(p, "body", "body")
        run = p.add_run()
        _글꼴_박기(run, self.kit.docx_font)
        _그림_넣기(run, plan.path, round(plan.width_mm * 36000), round(plan.height_mm * 36000))
        self._마지막이_표 = False

    # --- 가구 --------------------------------------------------------------------------------
    def _머리띠(self, plan: BandPlan) -> None:
        설정, 도장 = self.kit.furniture["band"], self.kit.furniture["stamp"]
        열폭 = [int(w) for w in 설정["cols"]]
        역할 = [("center", "band_left", "plain"), ("center", "band_teacher", "plain"),
                ("band_title", "band_title", "shade"), ("band_name", "band_name", "plain")]
        글들 = ["", "\n".join(plan.teacher_lines), plan.title, "\n".join(plan.name_lines)]
        if plan.stamp:
            # hwpx 의 도장은 머리띠 오른쪽 여백에 떠 있다(용지 기준). docx 는 떠 있는 도형을 쓰지 않으므로
            # 같은 줄의 표 칸으로 둔다: 머리띠 끝과 도장 왼쪽 사이 간격 = 도장 가로 위치 − (왼쪽 여백 + 머리띠 폭).
            간격 = max(0, int(도장["pos"]["horzOffset"]) - (self.r.page.left + int(설정["width"])))
            도장폭 = self.kit.body_width - int(설정["width"]) - 간격
            if 도장폭 <= 0:
                raise ValueError("머리띠 오른쪽에 확인도장 칸을 둘 자리가 없다 — band.width 와 stamp.pos 를 확인한다")
            if 간격:
                열폭.append(간격)
            열폭.append(도장폭)
        self._구분()
        표 = self._새_표(self.doc, 열폭, 1)
        행 = 표.rows[0]
        행.height_rule, 행.height = WD_ROW_HEIGHT_RULE.AT_LEAST, Twips(twip(int(설정["height"])))
        칸들 = 행.cells
        for i, ((문단_역할, 글자_역할, 상자), 글) in enumerate(zip(역할, 글들)):
            칸 = 칸들[i]
            tcPr = 칸._tc.get_or_add_tcPr()
            box = self.r.box[상자]
            self._테두리(tcPr, box)
            self._음영(tcPr, box.fill)
            칸.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            # 줄마다 문단이다(hwpx 머리띠 칸과 같다). 빈 줄에는 run 을 만들지 않는다.
            for j, 줄 in enumerate(글.split("\n")):
                문단 = 칸.paragraphs[0] if j == 0 else 칸.add_paragraph()
                self._문단_서식(문단, 문단_역할, 글자_역할)
                if 줄:
                    self._run(문단, 줄, 글자_역할)
        if plan.stamp:
            if len(칸들) == 6:
                간격_tcPr = 칸들[4]._tc.get_or_add_tcPr()
                self._테두리(간격_tcPr, None)
                # 구글 문서는 칸 여백(tcMar, 표 기본 좌우 102twip씩)보다 좁은 열을 만나면 열을
                # 넓혀 머리띠가 오른쪽 여백을 넘는다(실측 847→836px) — 이 칸만 좌우 여백을 0으로 막는다.
                간격_마진 = _두기(간격_tcPr, "w:tcMar", {}, _TCPR)
                _두기(간격_마진, "w:left", {"w:w": "0", "w:type": "dxa"}, _변_순서)
                _두기(간격_마진, "w:right", {"w:w": "0", "w:type": "dxa"}, _변_순서)
            도장칸 = 칸들[-1]
            tcPr = 도장칸._tc.get_or_add_tcPr()
            선 = Side("SOLID", _선폭_mm(str(도장["lineWidth"])), 도장["line"] or "#000000")
            self._테두리(tcPr, BoxStyle(선, 선, 선, 선, None))
            self._음영(tcPr, 도장["fill"])
            도장칸.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            문단 = 도장칸.paragraphs[0]
            self._문단_서식(문단, "center", "stamp")
            self._run(문단, 도장["text"], "stamp")
        self._마지막이_표 = True

    def _제목(self, plan: HeadingPlan) -> None:
        # plan.stamp 은 쓰지 않는다 — docx 의 확인도장은 머리띠의 칸이다.
        p = self.doc.add_paragraph()
        self._문단_서식(p, "body", "headline")
        번호 = self._run(p, 원문자(plan.number), "headline")
        번호.font.color.rgb = RGBColor.from_string(self.kit.furniture["circle"]["fill"].lstrip("#").upper())
        # 계획자는 None 이나 빈 문자열이 아닌 문구만 만든다(킷의 textbookFormat 은 비어 있지 않다)
        출처 = plan.textbook_ref is not None
        self._run(p, f" {plan.text} " if 출처 else f" {plan.text}", "headline")
        if 출처:
            self._run(p, plan.textbook_ref, "heading_ref")
        self._마지막이_표 = False

    def _키워드면(self, plan: KeywordPagePlan) -> None:
        설정 = self.kit.furniture["keywordPage"]
        self._구분()
        표 = self._새_표(self.doc, [int(설정["width"])], plan.rows)
        줄높이 = twip(int(설정["height"]) // plan.rows)
        # 줄(머리 칸 뺀 행)은 글이 없으니 계획한 높이 그대로(exact) 둔다. 최소 높이로 두면 빈 문단의
        # 줄 높이 + 칸 위아래 여백이 행을 키워 키워드면이 한 쪽을 넘는다 — 문단 줄도 표 기본 여백을 뺀
        # 칸 안쪽 높이로 누른다(행 높이를 최소값으로만 받는 구글 문서에서도 줄이 자라지 않는다).
        줄_안쪽 = 줄높이 - twip(_칸_여백["top"]) - twip(_칸_여백["bottom"])
        for i, 행 in enumerate(표.rows):
            행.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST if i == 0 else WD_ROW_HEIGHT_RULE.EXACTLY
            행.height = Twips(줄높이)
            상자 = "rule_head" if i == 0 else ("rule_last" if i == plan.rows - 1 else "rule_line")
            칸 = 행.cells[0]
            tcPr, box = 칸._tc.get_or_add_tcPr(), self.r.box[상자]
            self._테두리(tcPr, box)
            self._음영(tcPr, box.fill)
            문단 = 칸.paragraphs[0]
            글자 = "keyword_head" if i == 0 else "body"
            self._문단_서식(문단, "cell", 글자)
            if i == 0:
                self._run(문단, plan.head, "keyword_head")
            else:
                pf = 문단.paragraph_format
                pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.EXACTLY, Twips(줄_안쪽)
                # 구글 문서는 exact 줄간격이 없어 빈 문단을 문단부호 글자 크기(본문 10pt)로 재 줄이
                # 자란다 — 중첩표 앞 1pt 문단처럼 부호 크기를 1pt 로 누른다(실측 8줄 → 3줄 넘침).
                부호 = _두기(문단._p.get_or_add_pPr(), "w:rPr", {}, _PPR)
                _두기(부호, "w:sz", {"w:val": "2"}, _RPR)
                _두기(부호, "w:szCs", {"w:val": "2"}, _RPR)
                # LibreOffice 는 exact 행 높이에 칸 아래 여백을 더 얹는다(실측 356→384twip, 35줄이면 두
                # 줄이 다음 쪽으로) — 빈 칸이라 위아래 여백이 할 일이 없어 0 으로 둔다.
                마진 = _두기(tcPr, "w:tcMar", {}, _TCPR)
                _두기(마진, "w:top", {"w:w": "0", "w:type": "dxa"}, _변_순서)
                _두기(마진, "w:bottom", {"w:w": "0", "w:type": "dxa"}, _변_순서)
        self._마지막이_표 = True

    def draw(self, plan: Plan) -> None:
        if isinstance(plan, TablePlan):
            self._표(plan)
        elif isinstance(plan, ParagraphPlan):
            self._문단(plan)
        elif isinstance(plan, FigurePlan):
            self._그림(plan)
        elif isinstance(plan, BandPlan):
            self._머리띠(plan)
        elif isinstance(plan, HeadingPlan):
            self._제목(plan)
        elif isinstance(plan, KeywordPagePlan):
            self._키워드면(plan)
        else:
            raise TypeError(f"그릴 수 없는 계획: {type(plan).__name__}")

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
