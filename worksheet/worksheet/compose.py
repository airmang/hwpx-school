"""마크다운 + 킷 → 학습지(hwpx·docx)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from hwpx import HwpxDocument
from hwpx.experimental import render_layout_preview

from worksheet.backends.docx import DocxWriter
from worksheet.backends.hwpx import HwpxWriter
from worksheet.blocks import plan_node
from worksheet.checks import check_answer_labels, check_markdown, check_package
from worksheet.furniture import plan_band, plan_heading, plan_keyword_page
from worksheet.kit import Kit, effective_slots, load_kit
from worksheet.md import Heading, Sheet, parse_sheet
from worksheet.plan import BandPlan, HeadingPlan, KeywordPagePlan, Plan
from worksheet.kit_style import resolve_kit

_형식 = {".hwpx": "hwpx", ".docx": "docx"}


def output_format(path: Path) -> str:
    """출력 경로의 확장자(대소문자 무관)로 형식을 고른다 — 손으로 쓴 입력 오류는 한 문장으로."""
    확장자 = Path(path).suffix.lower()
    if 확장자 not in _형식:
        raise ValueError(f"출력 형식을 알 수 없다: {확장자 or '(없음)'} — .hwpx 또는 .docx")
    return _형식[확장자]


@dataclass(frozen=True)
class ComposeReport:
    pages_estimate: int | None  # hwpx 만 headless 근사값이 있다 — docx 는 None(실물 렌더로 본다)
    tables: int          # 문서 모델(HwpxDocument/Document)의 표 객체 수 — 레이아웃 쪽 나눔과 무관하다
    paragraphs: int      # 문서 모델(HwpxDocument/Document)의 문단 객체 수 — 레이아웃 쪽 나눔과 무관하다
    blocks: int
    headings: int
    validate_ok: bool


def build_plans(sheet: Sheet, kit: Kit, *, base_dir: Path) -> list[Plan]:
    """회차 하나의 조판 계획 목록 — 머리띠, (키워드면), 제목·블록을 문서 순서대로."""
    제목_있음 = any(isinstance(n, Heading) for n in sheet.nodes)
    계획들: list[Plan] = [
        plan_band(kit, title=sheet.title, slots=effective_slots(kit, grade=sheet.grade), stamp=제목_있음)
    ]
    if sheet.keyword_page:
        계획들.append(plan_keyword_page(kit))
    번호 = 0
    for node in sheet.nodes:
        if isinstance(node, Heading):
            번호 += 1
            계획들.append(plan_heading(kit, number=번호, text=node.text, textbook=node.textbook, with_stamp=(번호 == 1)))
        else:
            계획들.append(plan_node(kit, node, base_dir=base_dir))
    return 계획들


def compose(md_path: Path, kit_root: Path, out_path: Path) -> ComposeReport:
    md_path, out_path = Path(md_path), Path(out_path)
    # 무엇보다 먼저 본다 — 잘못된 출력 경로(모르는 확장자)는 md·킷을 읽기도, 파일을 하나
    # 만들기도 전에 거부한다("아무것도 하기 전에" 거부가 반쯤 만든 산출물을 없앤다).
    형식 = output_format(out_path)
    kit = load_kit(Path(kit_root))
    sheet = parse_sheet(md_path.read_text(encoding="utf-8"))
    if sheet.kit != kit.name:
        raise ValueError(f"킷 이름이 다르다: md={sheet.kit} 킷={kit.name}")

    # M-1·M-4 를 조판 전에 그대로 돌린다 — "M-1 초록 ⇒ 조판이 md 때문에 거부하지 않는다"가
    # 구조적으로 성립하게 한다. 여기서 걸러진 문제는 아래 계획자(plan_node)가 아무리 신중해도
    # 절대 늦게 터지지 않는다.
    문제 = check_markdown(sheet, kit, base_dir=md_path.parent) + check_answer_labels(sheet)
    if 문제:
        raise ValueError("조판할 수 없다:\n" + "\n".join(f"- {p}" for p in 문제))

    # 계획을 전부 만든 뒤에 그린다 — 계획 중 오류(그림 파일 등)가 나면 문서를 반쯤 만들지 않는다.
    계획들 = build_plans(sheet, kit, base_dir=md_path.parent)
    writer = HwpxWriter(kit) if 형식 == "hwpx" else DocxWriter(kit, resolve_kit(kit))
    for 계획 in 계획들:
        writer.draw(계획)
    writer.save(out_path)

    if 형식 == "docx":
        saved = Document(out_path)
        return ComposeReport(
            pages_estimate=None,
            tables=len(saved.tables),
            paragraphs=len(saved.paragraphs),
            blocks=sum(not isinstance(p, (BandPlan, HeadingPlan, KeywordPagePlan)) for p in 계획들),
            headings=sum(isinstance(p, HeadingPlan) for p in 계획들),
            validate_ok=check_package(out_path) == [],
        )

    saved = HwpxDocument.open(out_path)
    preview = render_layout_preview(out_path)
    return ComposeReport(
        pages_estimate=len(preview.pages),
        tables=len(saved.tables),
        paragraphs=len(saved.paragraphs),
        blocks=sum(not isinstance(p, (BandPlan, HeadingPlan, KeywordPagePlan)) for p in 계획들),
        headings=sum(isinstance(p, HeadingPlan) for p in 계획들),
        validate_ok=saved.validate().ok,
    )
