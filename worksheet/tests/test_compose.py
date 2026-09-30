import json
import re
import shutil
import zipfile

import pytest
from docx import Document
from hwpx.document import HwpxDocument

from worksheet.checks import check_markdown
from worksheet.cli import main as cli_main
from worksheet.compose import ComposeReport, compose, output_format
from worksheet.kit import load_kit
from worksheet.md import parse_sheet

MD = """---
kit: 시험킷
title: 탐색을 활용한 문제해결
grade: 2학년
keywordPage: false
---

## 탐색 알고리즘 종류
교과서: 33-35P

■ 정렬: 자료를 정해진 [[8]]에 따라 [[8]]대로 늘어놓는 일

:::비교표 cols="맹목적 탐색|휴리스틱 탐색" rows="정의|종류"
:::

## 맹목적 탐색

:::답칸 label="DFS 탐색 순서" lines=1
:::
"""


@pytest.fixture
def md파일(tmp_path):
    경로 = tmp_path / "회차03.md"
    경로.write_text(MD, encoding="utf-8")
    return 경로


def test_조판_보고가_구조를_센다(md파일, 킷_루트, tmp_path):
    out = tmp_path / "회차03.hwpx"
    보고 = compose(md파일, 킷_루트, out)

    assert isinstance(보고, ComposeReport)
    assert 보고.validate_ok is True
    assert 보고.headings == 2
    assert 보고.blocks == 3          # 정의빈칸 1 + 비교표 1 + 답칸 1
    assert 보고.tables == 4          # 머리띠 1 + 블록 3
    assert 보고.pages_estimate >= 1


def test_머리띠가_맨_앞이고_첫_제목에만_도장이_붙는다(md파일, 킷_루트, tmp_path):
    out = tmp_path / "회차03.hwpx"
    compose(md파일, 킷_루트, out)

    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert xml.index("탐색을 활용한 문제해결") < xml.index("탐색 알고리즘 종류")
    assert xml.count("확인도장") == 1
    assert xml.count("<hp:ellipse") == 2
    assert HwpxDocument.open(out).validate().ok


def test_키워드면_옵션이_표를_하나_더_만든다(md파일, 킷_루트, tmp_path):
    """쪽수로 단언하지 않는다 — 빠른 층의 쪽수는 실한컴과 다르다(근사값).

    키워드면이 실제로 한 쪽을 채우는지는 실한컴 렌더(M-7)로만 확인된다.
    """
    없음 = compose(md파일, 킷_루트, tmp_path / "a.hwpx")
    md파일.write_text(MD.replace("keywordPage: false", "keywordPage: true"), encoding="utf-8")
    있음 = compose(md파일, 킷_루트, tmp_path / "b.hwpx")

    assert 있음.tables == 없음.tables + 1
    assert 있음.validate_ok is True


# --- grade 는 죽은 입력이 아니다 — 회차(md)의 grade 가 킷의 grade 슬롯을 덮는다 -------------


def test_md의_grade가_킷의_grade_슬롯을_덮는다(md파일, 킷_루트, tmp_path):
    """킷_루트(conftest 시험_슬롯) 는 grade=2학년 — md가 3학년이면 산출물엔 3학년만 있어야 한다."""
    md파일.write_text(MD.replace("grade: 2학년", "grade: 3학년"), encoding="utf-8")
    out = tmp_path / "grade덮음.hwpx"
    compose(md파일, 킷_루트, out)

    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert "3학년" in xml
    assert "2학년" not in xml


def test_md에_grade가_없으면_킷_슬롯_그대로_쓴다(md파일, 킷_루트, tmp_path):
    md파일.write_text(MD.replace("grade: 2학년\n", ""), encoding="utf-8")
    out = tmp_path / "grade없음.hwpx"
    compose(md파일, 킷_루트, out)

    xml = zipfile.ZipFile(out).read("Contents/section0.xml").decode()
    assert "2학년" in xml  # 킷_루트 슬롯 자체가 2학년이다


def test_킷_이름이_어긋나면_거부한다(md파일, 킷_루트, tmp_path):
    md파일.write_text(MD.replace("kit: 시험킷", "kit: 다른킷"), encoding="utf-8")
    with pytest.raises(ValueError, match="킷 이름이 다르다"):
        compose(md파일, 킷_루트, tmp_path / "c.hwpx")


# --- B-2: "M-1 초록 ⇒ 조판이 md 때문에 거부하지 않는다"를 구조적으로 고정한다 -----------

_거짓_통과_본문 = {
    "cols_없는_비교표": '## 섹션\n\n:::비교표 rows="가"\n:::\n',
    "lines_abc": '## 섹션\n\n:::답칸 label="가" lines=abc\n:::\n',
    "lines_음수": '## 섹션\n\n:::답칸 label="가" lines=-1\n:::\n',
    "lines_0": '## 섹션\n\n:::답칸 label="가" lines=0\n:::\n',
    "빈_라벨설명": "## 섹션\n\n:::라벨설명\n:::\n",
    "없는_그림": "## 섹션\n\n![](없는파일.png)\n",
    "본문_있는_답칸": '## 섹션\n\n:::답칸 label="가" lines=1\n무시될 줄\n:::\n',
    "없는_칸_안_그림": '## 섹션\n\n:::나란히 cols="가|나"\n![](없는파일.png)|둘째\n:::\n',
    "칸_안_그림_형식_오류": '## 섹션\n\n:::나란히 cols="가|나"\n![나쁜형식|둘째\n:::\n',
}


@pytest.fixture
def 좁은_킷_루트(킷_루트, tmp_path):
    """'비교표'를 허용하지 않는 킷 — 킷의 허용 블록 대조(M-1) 전용 픽스처."""
    root = tmp_path / "좁은킷"
    shutil.copytree(킷_루트, root)
    데이터 = json.loads((root / "kit.json").read_text(encoding="utf-8"))
    데이터["blocks"] = [b for b in 데이터["blocks"] if b != "비교표"]
    (root / "kit.json").write_text(
        json.dumps(데이터, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return root


@pytest.mark.parametrize("이름", sorted(_거짓_통과_본문))
def test_거짓_통과_입력은_check_markdown과_compose_둘_다_잡는다(이름, 킷_루트, tmp_path):
    """완료 기준 3 — RED(지금은 []) → GREEN([기계] 검사와 compose가 같이 거부)."""
    md_본문 = "---\nkit: 시험킷\ntitle: t\n---\n\n" + _거짓_통과_본문[이름]
    md = tmp_path / "s.md"
    md.write_text(md_본문, encoding="utf-8")

    kit = load_kit(킷_루트)
    sheet = parse_sheet(md_본문)
    assert check_markdown(sheet, kit, base_dir=md.parent) != [], "① check_markdown이 비어 있으면 안 된다"

    with pytest.raises(ValueError, match="조판할 수 없다"):
        compose(md, 킷_루트, tmp_path / "out.hwpx")


def test_킷이_허용_안_하는_블록도_compose가_거부한다(좁은_킷_루트, tmp_path):
    md_본문 = (
        '---\nkit: 시험킷\ntitle: t\n---\n\n## 섹션\n\n:::비교표 cols="가|나" rows="다"\n:::\n'
    )
    md = tmp_path / "s.md"
    md.write_text(md_본문, encoding="utf-8")

    kit = load_kit(좁은_킷_루트)
    assert check_markdown(parse_sheet(md_본문), kit, base_dir=md.parent) != []

    with pytest.raises(ValueError, match="조판할 수 없다"):
        compose(md, 좁은_킷_루트, tmp_path / "out.hwpx")


def test_정상_md는_check_markdown이_비고_조판도_성공한다(md파일, 킷_루트, tmp_path):
    """거꾸로 — 기존 정상 md들은 ① [] ② 조판 성공."""
    kit = load_kit(킷_루트)
    sheet = parse_sheet(md파일.read_text(encoding="utf-8"))
    assert check_markdown(sheet, kit, base_dir=md파일.parent) == []

    보고 = compose(md파일, 킷_루트, tmp_path / "ok.hwpx")
    assert 보고.validate_ok is True


# --- "있지만 빈" 정수 속성은 거부가 아니라 기본값이다 ----------------------------------------
# lines=""·lines=" "·rows="" 를 그대로 두면 M-1은 []인데 compose()가 영문
# invalid literal for int()로 터진다. 판정: 빈 값은 기본값으로 읽는다(교사가 값을 지우고
# 속성 이름만 남긴 경우 — 거부보다 기본값이 덜 놀랍다). 그러므로 기대는 ① [] ② 조판 성공.

_빈_정수_속성_통과 = {
    "답칸_lines_빈문자열": '## 섹션\n\n:::답칸 label="가" lines=""\n:::\n',
    "답칸_lines_공백": '## 섹션\n\n:::답칸 label="가" lines=" "\n:::\n',
    "데이터표_rows_빈문자열": '## 섹션\n\n:::데이터표 head="가|나" rows=""\n:::\n',
}


@pytest.mark.parametrize("이름", sorted(_빈_정수_속성_통과))
def test_있지만_빈_정수_속성은_기본값으로_통과한다(이름, 킷_루트, tmp_path):
    md_본문 = "---\nkit: 시험킷\ntitle: t\n---\n\n" + _빈_정수_속성_통과[이름]
    md = tmp_path / "s.md"
    md.write_text(md_본문, encoding="utf-8")

    kit = load_kit(킷_루트)
    sheet = parse_sheet(md_본문)
    assert check_markdown(sheet, kit, base_dir=md.parent) == []

    보고 = compose(md, 킷_루트, tmp_path / "out.hwpx")
    assert 보고.validate_ok is True


# --- 모르는 속성은 M-1과 compose 둘 다 거부한다(오타 포함, 다른 블록에 쓴 labelWidth 포함) ---

_모르는_속성_거부_본문 = {
    "답칸_labelWidth": ':::답칸 label="가" labelWidth=3cm\n:::\n',
    "라벨설명_오타": ':::라벨설명 lable=3cm\n[1] 가\n: 나\n:::\n',
    "강조박스_오타": ':::강조박스 titel="x"\n본문\n:::\n',
}


@pytest.mark.parametrize("이름", sorted(_모르는_속성_거부_본문))
def test_모르는_속성이_있으면_check_markdown과_compose_둘_다_거부한다(이름, 킷_루트, tmp_path):
    md_본문 = "---\nkit: 시험킷\ntitle: t\n---\n\n## 섹션\n\n" + _모르는_속성_거부_본문[이름]
    md = tmp_path / "s.md"
    md.write_text(md_본문, encoding="utf-8")

    kit = load_kit(킷_루트)
    assert check_markdown(parse_sheet(md_본문), kit, base_dir=md.parent) != []

    with pytest.raises(ValueError, match="조판할 수 없다"):
        compose(md, 킷_루트, tmp_path / "out.hwpx")


def test_labelWidth가_너무_넓으면_compose도_거부한다(킷_루트, tmp_path):
    md_본문 = (
        "---\nkit: 시험킷\ntitle: t\n---\n\n## 섹션\n\n"
        ":::라벨설명 labelWidth=200cm\n[1] 가\n: 나\n:::\n"
    )
    md = tmp_path / "s.md"
    md.write_text(md_본문, encoding="utf-8")

    with pytest.raises(ValueError, match="조판할 수 없다"):
        compose(md, 킷_루트, tmp_path / "out.hwpx")


# --- 출력 형식 선택(확장자) -----------------------------------------------------------


@pytest.mark.parametrize("이름, 기대", [("a.hwpx", "hwpx"), ("a.docx", "docx"), ("a.DOCX", "docx"), ("a.HwpX", "hwpx")])
def test_출력_형식은_확장자로_대소문자_무관(이름, 기대, tmp_path):
    assert output_format(tmp_path / 이름) == 기대


@pytest.mark.parametrize("이름, 표기", [("a.pdf", ".pdf"), ("a", "(없음)")])
def test_모르는_확장자는_한_문장으로_거부(이름, 표기, tmp_path):
    with pytest.raises(ValueError, match=f"출력 형식을 알 수 없다: {re.escape(표기)} — .hwpx 또는 .docx"):
        output_format(tmp_path / 이름)


def test_같은_md를_docx로_조판한다(md파일, 킷_루트, tmp_path):
    보고 = compose(md파일, 킷_루트, tmp_path / "a.docx")
    assert 보고.validate_ok is True
    assert 보고.pages_estimate is None
    assert len(Document(tmp_path / "a.docx").tables) >= 1


def test_CLI_는_확장자_없는_출력을_종료코드_2로_거부(md파일, 킷_루트, tmp_path, capsys):
    코드 = cli_main(["compose", str(md파일), "--kit", str(킷_루트), "-o", str(tmp_path / "학습지")])
    assert 코드 == 2
    assert "오류: 출력 형식을 알 수 없다: (없음)" in capsys.readouterr().err
